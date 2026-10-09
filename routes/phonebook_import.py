# ╔══════════════════════════════════════════════════════════════╗
# ║                   phonebook_import.py                       ║
# ║  Импорт телефонного справочника из Excel                    ║
# ║                                                             ║
# ║  Маршруты:                                                  ║
# ║  GET  /phonebook/import/template  — скачать шаблон .xlsx    ║
# ║  POST /phonebook/import/upload    — загрузить файл (шаг 1)  ║
# ║  GET  /phonebook/import/resolve   — разрешить орг. (шаг 2)  ║
# ║  POST /phonebook/import/resolve   — сохранить маппинг       ║
# ║  POST /phonebook/import/confirm   — финальная загрузка (ш3) ║
# ╚══════════════════════════════════════════════════════════════╝

import os
import json
import uuid
from io import BytesIO
from datetime import datetime

from flask import (
    Blueprint, render_template, request,
    redirect, url_for, flash, send_file,
    session, jsonify
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db import get_db, UPLOADS_DIR
from core.auth_utils import login_required, admin_required
from core.activity_log import log_action

pb_import_bp = Blueprint('pb_import', __name__)

# ─── Константы ───────────────────────────────────────────────────────────────

TEMPLATE_HEADERS = [
    'Организация',
    'ФИО',
    'Должность',
    'Кабинет',
    'Рабочий тел.',
    'Доб.',
    'Личный тел.',
    'Email',
    'Примечания',
]

IMPORT_TEMP_DIR = os.path.join(UPLOADS_DIR, 'import_temp')
os.makedirs(IMPORT_TEMP_DIR, exist_ok=True)


# ─── Вспомогательные функции ─────────────────────────────────────────────────

def _get_all_orgs():
    conn = get_db()
    rows = conn.execute("SELECT id, name FROM phonebook_orgs ORDER BY name").fetchall()
    conn.close()
    return {r['name'].strip().lower(): r['id'] for r in rows}


def _parse_xlsx(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        return [], [f'Не удалось открыть файл: {e}']

    ws = wb.active
    errors = []
    rows = []

    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row and str(row[0]).strip() == 'Организация':
            header_row_idx = i
            break

    if header_row_idx is None:
        return [], ['Не найдена строка заголовков. Убедитесь что используете шаблон SONAR.']

    header_cells = list(ws.iter_rows(
        min_row=header_row_idx, max_row=header_row_idx, values_only=True
    ))[0]
    found_headers = [str(c).strip() if c else '' for c in header_cells[:len(TEMPLATE_HEADERS)]]
    if found_headers != TEMPLATE_HEADERS:
        errors.append(
            f'Заголовки колонок не совпадают с шаблоном. '
            f'Ожидалось: {TEMPLATE_HEADERS}. Найдено: {found_headers}'
        )
        return [], errors

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        cells = list(ws.iter_rows(
            min_row=row_idx, max_row=row_idx, values_only=True
        ))[0]

        if not any(cells[:len(TEMPLATE_HEADERS)]):
            continue

        full_name = str(cells[1]).strip() if cells[1] else ''
        if not full_name or full_name.lower() in ('фио', 'none', ''):
            errors.append(f'Строка {row_idx}: пустое поле «ФИО» — строка пропущена')
            continue

        rows.append({
            'org_name':       str(cells[0]).strip() if cells[0] else '',
            'full_name':      full_name,
            'position':       str(cells[2]).strip() if cells[2] else '',
            'room':           str(cells[3]).strip() if cells[3] else '',
            'phone_work':     str(cells[4]).strip() if cells[4] else '',
            'phone_ext':      str(cells[5]).strip() if cells[5] else '',
            'phone_personal': str(cells[6]).strip() if cells[6] else '',
            'email':          str(cells[7]).strip() if cells[7] else '',
            'notes':          str(cells[8]).strip() if cells[8] else '',
        })

    return rows, errors


def _save_import_session(data: dict) -> str:
    token = uuid.uuid4().hex
    path = os.path.join(IMPORT_TEMP_DIR, f'import_{token}.json')
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)
    return token


def _load_import_session(token: str):
    if not token:
        return None
    path = os.path.join(IMPORT_TEMP_DIR, f'import_{token}.json')
    if not os.path.exists(path):
        return None
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def _delete_import_session(token: str):
    path = os.path.join(IMPORT_TEMP_DIR, f'import_{token}.json')
    if os.path.exists(path):
        os.remove(path)


# ─── ШАГ 0: Скачать шаблон ───────────────────────────────────────────────────

@pb_import_bp.route('/phonebook/import/template')
@login_required
@admin_required
def import_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Справочник'

    hfill = PatternFill('solid', fgColor='1B5E7B')
    hfont = Font(bold=True, color='FFFFFF', size=10)
    efill = PatternFill('solid', fgColor='FFF9C4')
    br = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    ws.merge_cells(f'A1:{get_column_letter(len(TEMPLATE_HEADERS))}1')
    ws['A1'].value = 'Шаблон импорта телефонного справочника SONAR'
    ws['A1'].font = Font(bold=True, size=12, color='1B5E7B')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[1].height = 22

    ws.merge_cells(f'A2:{get_column_letter(len(TEMPLATE_HEADERS))}2')
    ws['A2'].value = (
        'Не изменяйте заголовки колонок (строка 3). '
        'Строка 4 — пример, удалите перед загрузкой. '
        'Поле «ФИО» обязательно.'
    )
    ws['A2'].font = Font(italic=True, size=9, color='888888')
    ws['A2'].alignment = Alignment(horizontal='center')

    col_widths = [30, 28, 24, 10, 16, 8, 16, 24, 30]
    for ci, (h, w) in enumerate(zip(TEMPLATE_HEADERS, col_widths), 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.fill = hfill
        c.font = hfont
        c.border = br
        c.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[3].height = 20

    example = [
        'Министерство инвестиций НО',
        'Иванов Иван Иванович',
        'Начальник отдела',
        '214',
        '+7 831 200-00-00',
        '101',
        '+7 912 000-00-00',
        'ivanov@example.ru',
        'Пример — удалите эту строку',
    ]
    for ci, val in enumerate(example, 1):
        c = ws.cell(row=4, column=ci, value=val)
        c.fill = efill
        c.border = br
        c.alignment = Alignment(vertical='center')
    ws.row_dimensions[4].height = 16

    ws.freeze_panes = 'A4'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f'sonar_phonebook_template_{datetime.now().strftime("%Y%m%d")}.xlsx'
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


# ─── ШАГ 1: Загрузка файла ───────────────────────────────────────────────────

@pb_import_bp.route('/phonebook/import/upload', methods=['POST'])
@login_required
@admin_required
def import_upload():
    f = request.files.get('import_file')
    if not f or not f.filename.endswith('.xlsx'):
        flash('Загрузите файл в формате .xlsx', 'error')
        return redirect(url_for('phonebook.phonebook'))

    tmp_name = f'upload_{uuid.uuid4().hex}.xlsx'
    tmp_path = os.path.join(IMPORT_TEMP_DIR, tmp_name)
    f.save(tmp_path)

    rows, errors = _parse_xlsx(tmp_path)
    os.remove(tmp_path)

    if errors and not rows:
        for e in errors:
            flash(e, 'error')
        return redirect(url_for('phonebook.phonebook'))

    if not rows:
        flash('Файл не содержит данных для импорта', 'error')
        return redirect(url_for('phonebook.phonebook'))

    known_orgs = _get_all_orgs()
    unknown_orgs = sorted(set(
        r['org_name'] for r in rows
        if r['org_name'] and r['org_name'].strip().lower() not in known_orgs
    ))

    token = _save_import_session({
        'rows': rows,
        'parse_errors': errors,
        'unknown_orgs': unknown_orgs,
    })
    session['import_token'] = token

    if unknown_orgs:
        return redirect(url_for('pb_import.import_resolve'))
    else:
        return redirect(url_for('pb_import.import_confirm_get'))


# ─── ШАГ 2: Разрешение организаций ───────────────────────────────────────────

@pb_import_bp.route('/phonebook/import/resolve', methods=['GET'])
@login_required
@admin_required
def import_resolve():
    token = session.get('import_token')
    data = _load_import_session(token)
    if not data:
        flash('Сессия импорта устарела. Загрузите файл заново.', 'error')
        return redirect(url_for('phonebook.phonebook'))

    conn = get_db()
    all_orgs = conn.execute('SELECT id, name FROM phonebook_orgs ORDER BY name').fetchall()
    conn.close()

    return render_template(
        'phonebook_import_resolve.html',
        unknown_orgs=data['unknown_orgs'],
        all_orgs=all_orgs,
        rows_count=len(data['rows']),
        parse_errors=data.get('parse_errors', []),
    )


@pb_import_bp.route('/phonebook/import/resolve', methods=['POST'])
@login_required
@admin_required
def import_resolve_post():
    token = session.get('import_token')
    data = _load_import_session(token)
    if not data:
        flash('Сессия импорта устарела. Загрузите файл заново.', 'error')
        return redirect(url_for('phonebook.phonebook'))

    mapping = {}
    for org_name in data['unknown_orgs']:
        slug = _org_slug(org_name)
        action = request.form.get(f'org_action_{slug}', 'create')
        map_id = request.form.get(f'org_map_{slug}', '')
        new_name = request.form.get(f'org_new_name_{slug}', org_name).strip() or org_name
        mapping[org_name] = {
            'action': action,
            'map_id': int(map_id) if map_id and map_id.isdigit() else None,
            'new_name': new_name,
        }

    data['org_mapping'] = mapping
    _save_import_session_by_token(token, data)
    return redirect(url_for('pb_import.import_confirm_get'))


# ─── ШАГ 3: Предпросмотр и финальная загрузка ────────────────────────────────

@pb_import_bp.route('/phonebook/import/confirm', methods=['GET'])
@login_required
@admin_required
def import_confirm_get():
    token = session.get('import_token')
    data = _load_import_session(token)
    if not data:
        flash('Сессия импорта устарела. Загрузите файл заново.', 'error')
        return redirect(url_for('phonebook.phonebook'))

    return render_template(
        'phonebook_import_resolve.html',
        unknown_orgs=[],
        all_orgs=[],
        rows_count=len(data['rows']),
        parse_errors=data.get('parse_errors', []),
        preview_mode=True,
        rows_preview=data['rows'][:10],
        update_duplicates=False,  # чекбокс всегда невыбран по умолчанию
    )


@pb_import_bp.route('/phonebook/import/confirm', methods=['POST'])
@login_required
@admin_required
def import_confirm_post():
    token = session.get('import_token')
    data = _load_import_session(token)
    if not data:
        flash('Сессия импорта устарела. Загрузите файл заново.', 'error')
        return redirect(url_for('phonebook.phonebook'))

    rows = data['rows']
    org_mapping = data.get('org_mapping', {})

    # ✔️ Главное исправление: читаем из POST шага 3 (чекбокс на странице предпросмотра)
    update_duplicates = request.form.get('update_duplicates') == '1'

    conn = get_db()
    known_orgs = {r['name'].strip().lower(): r['id']
                  for r in conn.execute('SELECT id, name FROM phonebook_orgs').fetchall()}

    created = 0
    updated = 0
    skipped = 0
    orgs_created = 0

    try:
        for orig_name, info in org_mapping.items():
            if info['action'] == 'create':
                new_name = info['new_name']
                key = new_name.strip().lower()
                if key not in known_orgs:
                    cur = conn.execute(
                        'INSERT INTO phonebook_orgs (name) VALUES (?)', (new_name,)
                    )
                    known_orgs[key] = cur.lastrowid
                    orgs_created += 1
            elif info['action'] == 'map' and info['map_id']:
                known_orgs[orig_name.strip().lower()] = info['map_id']

        for r in rows:
            org_name = r['org_name'].strip()
            org_key = org_name.lower()

            if org_name in org_mapping and org_mapping[org_name]['action'] == 'map':
                org_id = org_mapping[org_name]['map_id']
            else:
                org_id = known_orgs.get(org_key)

            existing = conn.execute(
                'SELECT id FROM phonebook WHERE full_name=? AND org_id IS ?',
                (r['full_name'], org_id)
            ).fetchone()

            if existing:
                if update_duplicates:
                    conn.execute("""
                        UPDATE phonebook SET
                            position      = ?,
                            room          = ?,
                            phone_work    = ?,
                            phone_ext     = ?,
                            phone_personal= ?,
                            email         = ?,
                            notes         = ?
                        WHERE id = ?
                    """, (
                        r['position'], r['room'], r['phone_work'],
                        r['phone_ext'], r['phone_personal'],
                        r['email'], r['notes'], existing['id'],
                    ))
                    updated += 1
                else:
                    skipped += 1
                continue

            conn.execute("""
                INSERT INTO phonebook
                    (org_id, full_name, position, room,
                     phone_work, phone_ext, phone_personal, email, notes)
                VALUES (?,?,?,?,?,?,?,?,?)
            """, (
                org_id, r['full_name'], r['position'], r['room'],
                r['phone_work'], r['phone_ext'], r['phone_personal'],
                r['email'], r['notes'],
            ))
            created += 1

        log_parts = [f'добавлено {created} контактов']
        if updated:  log_parts.append(f'обновлено {updated}')
        if skipped:  log_parts.append(f'пропущено {skipped} дублей')
        if orgs_created: log_parts.append(f'создано {orgs_created} орг.')

        log_action(conn, session['user_id'], 'create', None,
                   f'Импорт справочника: {", ".join(log_parts)}')
        conn.commit()

    except Exception as e:
        conn.rollback()
        conn.close()
        _delete_import_session(token)
        session.pop('import_token', None)
        flash(f'Ошибка при импорте: {e}', 'error')
        return redirect(url_for('phonebook.phonebook'))

    conn.close()
    _delete_import_session(token)
    session.pop('import_token', None)

    msg_parts = [f'добавлено {created} контактов']
    if updated:      msg_parts.append(f'обновлено {updated}')
    if skipped:      msg_parts.append(f'пропущено {skipped} дублей')
    if orgs_created: msg_parts.append(f'создано {orgs_created} новых организаций')

    flash(f'Импорт завершён: {", ".join(msg_parts)}', 'success')
    return redirect(url_for('phonebook.phonebook'))


# ─── AJAX: поиск организаций ────────────────────────────────────────────────

@pb_import_bp.route('/phonebook/import/orgs_search')
@login_required
def orgs_search():
    q = request.args.get('q', '').strip()
    conn = get_db()
    if q:
        rows = conn.execute(
            "SELECT id, name FROM phonebook_orgs WHERE name LIKE ? ORDER BY name LIMIT 30",
            (f'%{q}%',)
        ).fetchall()
    else:
        rows = conn.execute(
            'SELECT id, name FROM phonebook_orgs ORDER BY name LIMIT 50'
        ).fetchall()
    conn.close()
    return jsonify([{'id': r['id'], 'name': r['name']} for r in rows])


# ─── Утилиты ─────────────────────────────────────────────────────────────────

def _org_slug(name: str) -> str:
    import re
    return re.sub(r'[^a-zA-Z0-9а-яА-ЯёЁ]', '_', name)[:60]


def _save_import_session_by_token(token: str, data: dict):
    path = os.path.join(IMPORT_TEMP_DIR, f'import_{token}.json')
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)

# ─── Двухлистовый Excel: независимый сценарий ─────────────────────────────
# Дополнительные imports размещены здесь, чтобы старый импорт не изменять.
import secrets
import time
import zipfile
from hmac import compare_digest

from flask import abort
from services.phonebook_excel import (
    parse_phonebook_workbook, export_phonebook_workbook,
)
from services.phonebook_excel_plan import build_phonebook_plan
from services.phonebook_excel_apply import apply_phonebook_plan

PB_EXCEL_TTL = 3600
PB_EXCEL_MAX_BYTES = 20 * 1024 * 1024
PB_EXCEL_MAX_UNPACKED = 100 * 1024 * 1024


def _pb_excel_csrf_token():
    if not session.get('pb_excel_csrf'):
        session['pb_excel_csrf'] = secrets.token_hex(32)
    return session['pb_excel_csrf']


@pb_import_bp.app_context_processor
def _pb_excel_context():
    return {'pb_excel_csrf_token': _pb_excel_csrf_token}


def _pb_excel_check_csrf():
    expected = session.get('pb_excel_csrf', '')
    supplied = request.form.get('pb_excel_csrf', '')
    if not expected or not compare_digest(expected, supplied):
        abort(400, description='Защитный токен недействителен. Обновите страницу.')


def _pb_excel_path(token):
    if not isinstance(token, str) or len(token) != 32:
        raise ValueError('Некорректный пакет Excel')
    if any(c not in '0123456789abcdef' for c in token):
        raise ValueError('Некорректный пакет Excel')
    return os.path.join(IMPORT_TEMP_DIR, f'pb_excel_{token}.json')


def _pb_excel_save(token, data):
    path = _pb_excel_path(token)
    tmp = path + '.' + uuid.uuid4().hex + '.tmp'
    try:
        with open(tmp, 'x', encoding='utf-8') as handle:
            json.dump(data, handle, ensure_ascii=False)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


def _pb_excel_load():
    token = session.get('pb_excel_token')
    path = _pb_excel_path(token)
    if os.path.exists(path + '.claimed'):
        raise ValueError('Этот пакет уже использован. Загрузите книгу заново.')
    with open(path, encoding='utf-8') as handle:
        data = json.load(handle)
    if data['user_id'] != session.get('user_id'):
        raise ValueError('Пакет принадлежит другому пользователю')
    if not 0 <= time.time() - data['created_at'] <= PB_EXCEL_TTL:
        raise ValueError('Срок пакета истёк. Загрузите книгу заново.')
    return token, data


def _pb_excel_build(data):
    conn = get_db()
    try:
        conn.execute('BEGIN')
        return build_phonebook_plan(conn, data['parsed'], data['resolutions'])
    finally:
        conn.rollback()
        conn.close()


def _pb_excel_error(exc):
    flash(f'Excel-справочник: {exc}', 'error')
    return redirect(url_for('phonebook.phonebook'))


@pb_import_bp.route('/phonebook/excel/export')
@login_required
@admin_required
def excel_export():
    conn = get_db()
    try:
        conn.execute('BEGIN')
        buf = export_phonebook_workbook(conn)
    finally:
        conn.rollback()
        conn.close()
    return send_file(
        buf, as_attachment=True,
        download_name=f'phonebook_edit_{datetime.now():%Y%m%d_%H%M%S}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@pb_import_bp.route('/phonebook/excel/upload', methods=['POST'])
@login_required
@admin_required
def excel_upload():
    _pb_excel_check_csrf()
    upload = request.files.get('import_file')
    if not upload or not (upload.filename or '').lower().endswith('.xlsx'):
        return _pb_excel_error('Выберите файл .xlsx')
    raw = upload.stream.read(PB_EXCEL_MAX_BYTES + 1)
    if len(raw) > PB_EXCEL_MAX_BYTES:
        return _pb_excel_error('Максимальный размер книги — 20 МБ')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > PB_EXCEL_MAX_UNPACKED:
                raise ValueError('Распакованная книга превышает 100 МБ')
        parsed = parse_phonebook_workbook(BytesIO(raw))
        token = uuid.uuid4().hex
        data = {
            'user_id': session['user_id'], 'created_at': time.time(),
            'parsed': parsed, 'resolutions': {},
        }
        _pb_excel_save(token, data)
        session['pb_excel_token'] = token
    except Exception as exc:
        return _pb_excel_error(exc)
    return redirect(url_for('pb_import.excel_preview'))


@pb_import_bp.route('/phonebook/excel/preview')
@login_required
@admin_required
def excel_preview():
    try:
        token, data = _pb_excel_load()
        plan = _pb_excel_build(data)
        data['approved_plan'] = plan
        data['preview_version'] = uuid.uuid4().hex
        _pb_excel_save(token, data)
    except Exception as exc:
        return _pb_excel_error(exc)
    return render_template(
        'phonebook_excel_preview.html', plan=plan,
        preview_version=data['preview_version'],
        resolutions=data['resolutions'],
    )


@pb_import_bp.route('/phonebook/excel/resolve', methods=['POST'])
@login_required
@admin_required
def excel_resolve():
    _pb_excel_check_csrf()
    try:
        token, data = _pb_excel_load()
        if request.form.get('preview_version') != data.get('preview_version'):
            raise ValueError('Предпросмотр изменился. Обновите страницу.')
        for conflict in data['approved_plan'].get('conflicts', []):
            key = conflict['key']
            choice = request.form.get('resolve_' + key, '')
            if not choice:
                continue
            if choice == 'create' and conflict['allow_create']:
                data['resolutions'][key] = {'create': True}
            elif choice.isdigit() and int(choice) in conflict['candidates']:
                data['resolutions'][key] = {'id': int(choice)}
            else:
                raise ValueError('Недопустимое решение совпадения')
        data.pop('approved_plan', None)
        data.pop('preview_version', None)
        _pb_excel_save(token, data)
    except Exception as exc:
        return _pb_excel_error(exc)
    return redirect(url_for('pb_import.excel_preview'))


@pb_import_bp.route('/phonebook/excel/apply', methods=['POST'])
@login_required
@admin_required
def excel_apply():
    _pb_excel_check_csrf()
    if request.form.get('confirm_apply') != '1':
        return _pb_excel_error('Нужно подтвердить применение изменений')
    try:
        token, data = _pb_excel_load()
        if request.form.get('preview_version') != data.get('preview_version'):
            raise ValueError('Предпросмотр изменился. Обновите страницу.')
        approved = data.get('approved_plan', {})
        if not approved.get('can_apply'):
            raise ValueError('Сначала исправьте ошибки и совпадения')
        path = _pb_excel_path(token)
        # O_EXCL: только один запрос может использовать пакет, даже после сбоя.
        fd = os.open(path + '.claimed', os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        os.close(fd)
    except Exception as exc:
        return _pb_excel_error(exc)
    try:
        # После захвата проверяем, что другой запрос не заменил предпросмотр.
        with open(path, encoding='utf-8') as handle:
            current = json.load(handle)
        if current != data:
            raise ValueError('Пакет изменился. Загрузите книгу заново.')
        conn = get_db()
        try:
            counts = apply_phonebook_plan(
                conn, data['parsed'], approved, session['user_id'],
                log_action, data['resolutions'],
            )
        finally:
            conn.close()
        labels = {'organizations': 'организации', 'contacts': 'контакты'}
        parts = [f"{labels[k]}: добавлено {v['create']}, "
                 f"изменено {v['update']}, удалено {v['delete']}"
                 for k, v in counts.items()]
        flash('Excel применён: ' + '; '.join(parts), 'success')
    except Exception as exc:
        flash(f'Excel не применён: {exc}. Загрузите книгу заново.', 'error')
    finally:
        session.pop('pb_excel_token', None)
        try:
            os.remove(path)
        except OSError:
            pass
    return redirect(url_for('phonebook.phonebook'))

