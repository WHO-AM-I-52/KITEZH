"""Read-only Excel export for quarterly InvestMap data updates."""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


_STATUS_LABELS = {
    "planned": "Не направлено",
    "request_sent": "Запрос направлен",
    "response_received": "Получен ответ",
    "overdue": "Требуется уточнение",
}

_SOURCE_TYPE_LABELS = {
    "district": "Район",
    "organization": "Организация",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
_TOTAL_FILL = PatternFill("solid", fgColor="E2F0D9")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_BOLD_FONT = Font(bold=True)

_RECORD_HEADERS = [
    "№",
    "Категория",
    "Код категории",
    "Получатель",
    "Тип получателя",
    "Статус",
    "Исходящий № запроса",
    "Входящий № ответа",
    "Количество изменений",
    "Запросов",
    "Ответов",
    "Иных документов",
    "Комментарий",
    "Дата обновления",
    "Обновил",
]

_CATEGORY_HEADERS = [
    "№",
    "Получатель",
    "Тип получателя",
    "Статус",
    "Исходящий № запроса",
    "Входящий № ответа",
    "Количество изменений",
    "Запросов",
    "Ответов",
    "Иных документов",
    "Комментарий",
    "Дата обновления",
    "Обновил",
]

_SUMMARY_HEADERS = [
    "Категория",
    "Всего строк",
    "Не направлено",
    "Запрос направлен",
    "Получен ответ",
    "Требуется уточнение",
    "Строк с изменениями",
    "Всего изменений",
    "Запросов",
    "Ответов",
    "Иных документов",
    "Готовность, %",
]

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _display_value(value: Any, default: str = "—") -> Any:
    """Возвращает значение для Excel, заменяя пустые строки на default."""
    if value is None:
        return default

    if isinstance(value, str):
        value = value.strip()
        return value or default

    return value


def _status_label(status: Any) -> str:
    status = str(status or "").strip()
    return _STATUS_LABELS.get(status, status or "—")


def _source_type_label(source_type: Any) -> str:
    source_type = str(source_type or "").strip()
    return _SOURCE_TYPE_LABELS.get(source_type, source_type or "—")


def _int_value(value: Any) -> int:
    """Безопасно приводит счётчик из sqlite.Row к неотрицательному int."""
    try:
        return max(int(value or 0), 0)
    except (TypeError, ValueError):
        return 0


def _period_label(plan: Any) -> str:
    start = str(plan["period_start"]).replace("-", ".")
    end = str(plan["period_end"]).replace("-", ".")
    return f"{start} — {end}"


def _sheet_title(
    raw_title: str,
    used_titles: set[str],
) -> str:
    """Создаёт уникальное имя листа Excel длиной не более 31 символа."""
    clean_title = _INVALID_SHEET_CHARS.sub(" ", raw_title)
    clean_title = " ".join(clean_title.split()) or "Категория"
    base_title = clean_title[:31]
    title = base_title
    suffix = 2

    while title in used_titles:
        suffix_text = f" ({suffix})"
        title = f"{base_title[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1

    used_titles.add(title)
    return title


def _style_header(
    sheet: Any,
    row_number: int,
    headers: list[str],
) -> None:
    for column_number, header in enumerate(headers, start=1):
        cell = sheet.cell(
            row=row_number,
            column=column_number,
            value=header,
        )
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.freeze_panes = f"A{row_number + 1}"


def _apply_widths(sheet: Any, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _style_title(
    sheet: Any,
    title: str,
    subtitle: str,
    columns_count: int,
) -> None:
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=columns_count,
    )
    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="left")

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=columns_count,
    )
    subtitle_cell = sheet.cell(row=2, column=1, value=subtitle)
    subtitle_cell.font = Font(italic=True, color="666666")
    subtitle_cell.alignment = Alignment(horizontal="left")


def _record_values(
    record: Any,
    number: int,
    *,
    include_category: bool,
) -> list[Any]:
    common_values = [
        number,
        _display_value(record["source_name_snapshot"]),
        _source_type_label(record["source_type"]),
        _status_label(record["status"]),
        _display_value(record["request_letter_number"]),
        _display_value(record["response_letter_number"]),
        _int_value(record["changes_count"]),
        _int_value(record["request_documents_count"]),
        _int_value(record["response_documents_count"]),
        _int_value(record["other_documents_count"]),
        _display_value(record["note"]),
        _display_value(record["updated_at"]),
        _display_value(record["updated_by_name"]),
    ]

    if not include_category:
        return common_values

    return [
        number,
        _display_value(record["category_name"]),
        _display_value(record["category_code"]),
        *common_values[1:],
    ]


def _apply_record_row_style(
    sheet: Any,
    row_number: int,
    columns_count: int,
) -> None:
    for column_number in range(1, columns_count + 1):
        cell = sheet.cell(row=row_number, column=column_number)
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


def _append_summary_sheet(
    workbook: Workbook,
    plan: Any,
    records: list[Any],
) -> None:
    sheet = workbook.active
    sheet.title = "Сводка"

    _style_title(
        sheet,
        "Итоги квартальной актуализации Инвесткарты",
        f"Период: {_period_label(plan)}",
        len(_SUMMARY_HEADERS),
    )
    _style_header(sheet, 4, _SUMMARY_HEADERS)

    grouped: dict[int, list[Any]] = defaultdict(list)
    for record in records:
        grouped[int(record["category_id"])].append(record)

    total_rows = 0
    total_statuses = defaultdict(int)
    total_changed_records = 0
    total_changes = 0
    total_request_documents = 0
    total_response_documents = 0
    total_other_documents = 0

    for category_records in grouped.values():
        first_record = category_records[0]
        status_counts = defaultdict(int)

        changes_records_count = 0
        changes_total = 0
        request_documents_total = 0
        response_documents_total = 0
        other_documents_total = 0

        for record in category_records:
            status = str(record["status"] or "")
            status_counts[status] += 1

            changes_count = _int_value(record["changes_count"])
            if changes_count > 0:
                changes_records_count += 1
            changes_total += changes_count

            request_documents_total += _int_value(
                record["request_documents_count"]
            )
            response_documents_total += _int_value(
                record["response_documents_count"]
            )
            other_documents_total += _int_value(
                record["other_documents_count"]
            )

        records_count = len(category_records)
        received_count = status_counts["response_received"]
        readiness_percent = (
            received_count / records_count
            if records_count
            else 0
        )

        sheet.append(
            [
                _display_value(first_record["category_name"]),
                records_count,
                status_counts["planned"],
                status_counts["request_sent"],
                status_counts["response_received"],
                status_counts["overdue"],
                changes_records_count,
                changes_total,
                request_documents_total,
                response_documents_total,
                other_documents_total,
                readiness_percent,
            ]
        )

        total_rows += records_count
        for status, count in status_counts.items():
            total_statuses[status] += count
        total_changed_records += changes_records_count
        total_changes += changes_total
        total_request_documents += request_documents_total
        total_response_documents += response_documents_total
        total_other_documents += other_documents_total

    total_readiness_percent = (
        total_statuses["response_received"] / total_rows
        if total_rows
        else 0
    )

    total_row = sheet.max_row + 1
    sheet.append(
        [
            "Итого",
            total_rows,
            total_statuses["planned"],
            total_statuses["request_sent"],
            total_statuses["response_received"],
            total_statuses["overdue"],
            total_changed_records,
            total_changes,
            total_request_documents,
            total_response_documents,
            total_other_documents,
            total_readiness_percent,
        ]
    )

    for cell in sheet[total_row]:
        cell.fill = _TOTAL_FILL
        cell.font = _BOLD_FONT

    for cell in sheet["L"][4:]:
        cell.number_format = "0.0%"

    sheet.auto_filter.ref = (
        f"A4:{get_column_letter(len(_SUMMARY_HEADERS))}{total_row}"
    )
    _apply_widths(
        sheet,
        [34, 15, 17, 18, 17, 23, 20, 18, 13, 13, 18, 16],
    )


def _append_records_sheet(
    workbook: Workbook,
    title: str,
    subtitle: str,
    records: list[Any],
    *,
    include_category: bool,
    used_titles: set[str],
) -> None:
    sheet = workbook.create_sheet(_sheet_title(title, used_titles))
    headers = _RECORD_HEADERS if include_category else _CATEGORY_HEADERS

    _style_title(
        sheet,
        title,
        subtitle,
        len(headers),
    )
    _style_header(sheet, 4, headers)

    for number, record in enumerate(records, start=1):
        sheet.append(
            _record_values(
                record,
                number,
                include_category=include_category,
            )
        )
        _apply_record_row_style(
            sheet,
            sheet.max_row,
            len(headers),
        )

    last_row = max(sheet.max_row, 4)
    sheet.auto_filter.ref = (
        f"A4:{get_column_letter(len(headers))}{last_row}"
    )

    if include_category:
        _apply_widths(
            sheet,
            [
                7,
                32,
                16,
                34,
                18,
                22,
                22,
                22,
                19,
                12,
                12,
                18,
                38,
                20,
                28,
            ],
        )
    else:
        _apply_widths(
            sheet,
            [
                7,
                34,
                18,
                22,
                22,
                22,
                19,
                12,
                12,
                18,
                38,
                20,
                28,
            ],
        )


def build_update_plan_export_xlsx(
    plan: Any,
    records: list[Any],
) -> BytesIO:
    """
    Формирует read-only XLSX-итоги квартального плана актуализации.

    Создаёт листы:
    - Сводка;
    - Все изменения (только changes_count > 0);
    - по одному листу на категорию;
    - Без изменений (changes_count = 0).

    Функция не изменяет базу данных и не работает с файловым хранилищем.
    """
    workbook = Workbook()
    used_titles = {"Сводка"}

    _append_summary_sheet(workbook, plan, records)

    period = _period_label(plan)
    changed_records = [
        record
        for record in records
        if _int_value(record["changes_count"]) > 0
    ]
    unchanged_records = [
        record
        for record in records
        if _int_value(record["changes_count"]) == 0
    ]

    _append_records_sheet(
        workbook,
        "Все изменения",
        f"Период: {period}. Только строки с количеством изменений больше нуля.",
        changed_records,
        include_category=True,
        used_titles=used_titles,
    )

    category_groups: dict[int, list[Any]] = defaultdict(list)
    for record in records:
        category_groups[int(record["category_id"])].append(record)

    ordered_groups = sorted(
        category_groups.values(),
        key=lambda group: (
            _int_value(group[0]["category_sort_order"]),
            str(group[0]["category_name"] or "").casefold(),
            int(group[0]["category_id"]),
        ),
    )

    for category_records in ordered_groups:
        category_name = str(
            category_records[0]["category_name"] or "Категория"
        )
        _append_records_sheet(
            workbook,
            category_name,
            f"Период: {period}. Все строки категории.",
            category_records,
            include_category=False,
            used_titles=used_titles,
        )

    _append_records_sheet(
        workbook,
        "Без изменений",
        f"Период: {period}. Строки с количеством изменений, равным нулю.",
        unchanged_records,
        include_category=True,
        used_titles=used_titles,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
